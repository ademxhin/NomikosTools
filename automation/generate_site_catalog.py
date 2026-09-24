from __future__ import annotations

import json
import re
import unicodedata
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "outputs" / "nomikos-hourly-sync" / "Nomikos Tools - Ενημερωμένο.xlsx"
PLACEHOLDER_IMAGE = "https://images.unsplash.com/photo-1504148455328-c376907d081c?auto=format&fit=crop&w=800&q=80"

GREEK_LATIN = str.maketrans({
    "α": "a", "β": "v", "γ": "g", "δ": "d", "ε": "e", "ζ": "z", "η": "i", "θ": "th",
    "ι": "i", "κ": "k", "λ": "l", "μ": "m", "ν": "n", "ξ": "x", "ο": "o", "π": "p",
    "ρ": "r", "σ": "s", "ς": "s", "τ": "t", "υ": "y", "φ": "f", "χ": "ch", "ψ": "ps", "ω": "o",
})


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    value = str(value).strip()
    return "" if value in {"", "-"} else value


def number(value: object, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return default


def slug(value: str, fallback: str) -> str:
    normalized = unicodedata.normalize("NFD", value.lower()).translate(GREEK_LATIN)
    normalized = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return normalized or fallback


def unique_slug(value: str, used: set[str], fallback: str) -> str:
    base = slug(value, fallback)
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}-{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def find_product_sheet(workbook):
    for sheet in workbook.worksheets:
        headers = {text(cell.value) for cell in sheet[1]}
        if "UID" in headers and "Όνομα" in headers:
            return sheet
    raise ValueError("Δεν βρέθηκε φύλλο προϊόντων με UID και Όνομα.")


def load_rows(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = find_product_sheet(workbook)
    headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
    headers["EAN"] = headers.get("EAN", headers.get("ΕΑΝ"))
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        uid = text(values[headers["UID"] - 1])
        name = text(values[headers["Όνομα"] - 1])
        if not uid or not name or not uid.isdigit():
            continue
        def get(label: str) -> object:
            column = headers.get(label)
            return values[column - 1] if column else None
        rows.append({
            "uid": uid,
            "sku_id": text(get("SKU ID")),
            "name": name,
            "category": text(get("Κατηγορία προϊόντος")) or "Λοιπά προϊόντα",
            "manufacturer": text(get("Κατασκευαστής")) or "Άγνωστος κατασκευαστής",
            "ean": text(get("EAN")),
            "mpn": text(get("MPN")),
            "availability": text(get("Διαθεσιμότητα")),
            "quantity": number(get("Ποσότητα")),
            "retail_price": number(get("Τιμή με ΦΠΑ")),
            "cost_price": number(get("Τιμή χονδρικής")),
            "vat_rate": number(get("ΦΠΑ %"), 24) / 100,
            "image": text(get("XML Εικόνα")),
            "supplier_code": text(get("XML Κωδικός")),
            "supplier_stock": text(get("XML Διαθεσιμότητα")),
            "weight": number(get("XML Βάρος")),
        })
    if not rows:
        raise ValueError("Δεν βρέθηκαν έγκυρα προϊόντα στο Excel.")
    return rows


def make_catalog(rows: list[dict[str, object]]) -> tuple[list[dict], list[dict], list[dict]]:
    category_slugs: dict[str, str] = {}
    brand_slugs: dict[str, str] = {}
    used_category_slugs: set[str] = set()
    used_brand_slugs: set[str] = set()
    for index, row in enumerate(rows, start=1):
        category = str(row["category"])
        manufacturer = str(row["manufacturer"])
        category_slugs.setdefault(category, unique_slug(category, used_category_slugs, f"category-{index}"))
        brand_slugs.setdefault(manufacturer, unique_slug(manufacturer, used_brand_slugs, f"brand-{index}"))

    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    categories: list[dict] = []
    for index, (name, category_slug) in enumerate(category_slugs.items(), start=1):
        category_rows = [row for row in rows if row["category"] == name]
        image = next((str(row["image"]) for row in category_rows if row["image"]), PLACEHOLDER_IMAGE)
        categories.append({
            "id": f"cat-{category_slug}", "name": name, "slug": category_slug,
            "description": f"Προϊόντα στην κατηγορία {name}.", "image": image,
            "iconName": "Boxes", "level": 1, "order": index, "active": True,
            "featured": index <= 12,
        })

    brands: list[dict] = []
    for index, (name, brand_slug) in enumerate(brand_slugs.items(), start=1):
        brand_rows = [row for row in rows if row["manufacturer"] == name]
        logo = next((str(row["image"]) for row in brand_rows if row["image"]), PLACEHOLDER_IMAGE)
        brands.append({
            "id": f"brand-{brand_slug}", "name": name, "slug": brand_slug,
            "logo": logo, "description": f"Προϊόντα {name} από τον κατάλογο Excel.",
            "country": "", "active": True, "featured": index <= 12,
        })

    products: list[dict] = []
    for row in rows:
        quantity = max(0, int(row["quantity"]))
        availability_text = f"{row['availability']} {row['supplier_stock']}".casefold()
        if quantity > 0:
            status = "in_stock"
        elif "ημέρ" in availability_text or "ετοιμάζεται" in availability_text or "παραγγελ" in availability_text:
            status = "on_order"
        else:
            status = "out_of_stock"
        product_slug = unique_slug(str(row["name"]), set(p["slug"] for p in products), f"product-{row['uid']}")
        category_id = f"cat-{category_slugs[str(row['category'])]}"
        brand_id = f"brand-{brand_slugs[str(row['manufacturer'])]}"
        sku = str(row["mpn"] or row["sku_id"] or row["uid"])
        image = str(row["image"] or PLACEHOLDER_IMAGE)
        product = OrderedDict([
            ("id", f"prod-{row['uid']}"),
            ("name", row["name"]),
            ("slug", product_slug),
            ("shortDescription", f"{row['manufacturer']} · {row['category']} · Κωδικός {sku}"),
            ("fullDescription", f"{row['name']}. Κατασκευαστής: {row['manufacturer']}. Κατηγορία: {row['category']}."),
            ("sku", sku), ("ean", row["ean"]), ("manufacturerCode", row["mpn"] or sku),
            ("brandId", brand_id), ("categoryIds", [category_id]),
            ("price", round(float(row["retail_price"]), 2)),
            ("costPrice", round(float(row["cost_price"]), 4)),
            ("vatRate", round(float(row["vat_rate"]), 4)),
            ("stockQuantity", quantity), ("lowStockThreshold", 1),
            ("availabilityStatus", status), ("unit", "Τεμάχιο"), ("packageQuantity", 1),
            ("images", [image]), ("attributes", {"manufacturer": row["manufacturer"], "category": row["category"], "supplierCode": row["supplier_code"], "supplierStock": row["supplier_stock"]}),
            ("specifications", [{"key": "Κατασκευαστής", "value": row["manufacturer"]}, {"key": "Κατηγορία", "value": row["category"]}, {"key": "Διαθεσιμότητα προμηθευτή", "value": row["supplier_stock"] or row["availability"]}]),
            ("includedItems", []), ("compatibleProductIds", []), ("relatedProductIds", []), ("manualFiles", []),
            ("warranty", ""), ("weight", round(float(row["weight"]), 3)),
            ("dimensions", {"length": 0, "width": 0, "height": 0}),
            ("seoTitle", f"{row['name']} | Nomikos Tools"),
            ("seoDescription", f"{row['name']} από {row['manufacturer']}."),
            ("active", True), ("featured", False), ("popular", False),
            ("rating", 0), ("reviewCount", 0), ("createdAt", timestamp), ("updatedAt", timestamp),
        ])
        products.append(product)
    return products, categories, brands


def write_ts(path: Path, export_name: str, import_name: str, records: list[dict]) -> None:
    path.write_text(
        f"import {{ {import_name} }} from '../types/{'product' if import_name == 'Product' else import_name.lower()}';\n\n"
        f"export const {export_name}: {import_name}[] = "
        + json.dumps(records, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    rows = load_rows(source)
    products, categories, brands = make_catalog(rows)
    write_ts(ROOT / "src" / "data" / "productsData.ts", "initialProducts", "Product", products)
    write_ts(ROOT / "src" / "data" / "categoriesData.ts", "initialCategories", "Category", categories)
    write_ts(ROOT / "src" / "data" / "brandsData.ts", "initialBrands", "Brand", brands)
    print(json.dumps({"products": len(products), "categories": len(categories), "brands": len(brands)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
