from __future__ import annotations

import argparse
import csv
import logging
from logging.handlers import RotatingFileHandler
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
import xml.etree.ElementTree as ET
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


APP_DIR = Path(__file__).resolve().parent
DEFAULT_ENV_FILE = APP_DIR / ".env"
DEFAULT_OUTPUT_DIR = APP_DIR.parent / "outputs" / "nomikos-hourly-sync"
DEFAULT_LOG_DIR = APP_DIR / "logs"
DEFAULT_STATE_DIR = APP_DIR / "state"

PRODUCT_SHEET = "Προϊόντα που εμφανίζονται"
SUMMARY_SHEET = "Έλεγχος XML"
OUTPUT_FILENAME = "Nomikos Tools - Ενημερωμένο.xlsx"
XML_PRODUCTS_FILENAME = "XML Products - Πλήρης Κατάλογος.xlsx"
PRODUCT_AUDIT_FILENAME = "Έλεγχος προϊόντων Skroutz vs XML.xlsx"

EXTRA_HEADERS = [
    "XML Κατάσταση Αντιστοίχισης",
    "XML Κωδικός",
    "XML Διαθεσιμότητα",
    "XML Τιμή Χονδρικής χωρίς ΦΠΑ",
    "XML Τιμή Χονδρικής με ΦΠΑ",
    "XML Βάρος",
    "XML Εικόνα",
    "Τελευταίος Συγχρονισμός",
    "Έλεγχος",
]

XML_EXPORT_HEADERS = [
    "XML Κωδικός",
    "Alternativecode / MPN",
    "Barcode / EAN",
    "Τίτλος",
    "Διαθεσιμότητα",
    "Τιμή Χονδρικής χωρίς ΦΠΑ",
    "Τιμή Χονδρικής με ΦΠΑ",
    "Βάρος",
    "Εικόνα",
]

PRODUCT_AUDIT_HEADERS = [
    "UID",
    "SKU ID",
    "Όνομα",
    "Κατηγορία προϊόντος",
    "Κατασκευαστής",
    "EAN",
    "MPN",
    "Skroutz URL",
    "Κατάσταση",
    "XML Κωδικός",
    "XML Διαθεσιμότητα",
    "XML Εικόνα",
]


@dataclass(frozen=True)
class ProductRecord:
    code: str
    alternative_code: str
    barcode: str
    title: str
    stock: str
    wholesale_no_vat: Decimal | None
    wholesale_with_vat: Decimal | None
    weight: Decimal | None
    image: str


@dataclass(frozen=True)
class MatchResult:
    status: str
    product: ProductRecord | None


def normalize(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text in {"", "-"} else text


def decimal_or_none(value: object) -> Decimal | None:
    text = normalize(value).replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key.strip(), value)


def configure_logging(log_dir: Path, verbose: bool) -> logging.Logger:
    log_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("nomikos_sync")
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    file_handler = RotatingFileHandler(
        log_dir / "sync.log", maxBytes=2_000_000, backupCount=5, encoding="utf-8"
    )
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except OSError:
            pass
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(formatter)
    logger.addHandler(console)
    return logger


class RunLock:
    def __init__(self, lock_path: Path, stale_after_seconds: int = 3600) -> None:
        self.lock_path = lock_path
        self.stale_after_seconds = stale_after_seconds
        self.acquired = False

    def __enter__(self) -> "RunLock":
        self.lock_path.parent.mkdir(parents=True, exist_ok=True)
        if self.lock_path.exists():
            age = time.time() - self.lock_path.stat().st_mtime
            if age > self.stale_after_seconds:
                self.lock_path.unlink(missing_ok=True)
        try:
            descriptor = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as exc:
            raise RuntimeError("Υπάρχει ήδη συγχρονισμός σε εξέλιξη.") from exc
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(f"pid={os.getpid()}\nstarted={datetime.now().astimezone().isoformat()}\n")
        self.acquired = True
        return self

    def __exit__(self, exc_type, exc, traceback) -> None:
        if self.acquired:
            self.lock_path.unlink(missing_ok=True)


def download_xml(url: str, destination: Path, timeout: int, logger: logging.Logger) -> None:
    request = Request(url, headers={"User-Agent": "NomikosToolsSync/1.0"})
    last_error: Exception | None = None
    for attempt in range(1, 4):
        try:
            logger.info("Λήψη XML, προσπάθεια %s/3", attempt)
            with urlopen(request, timeout=timeout) as response, destination.open("wb") as output:
                content_type = response.headers.get("Content-Type", "")
                if "xml" not in content_type.lower() and "octet-stream" not in content_type.lower():
                    logger.warning("Μη αναμενόμενο Content-Type: %s", content_type)
                shutil.copyfileobj(response, output)
            if destination.stat().st_size < 1_000:
                raise ValueError("Το XML που λήφθηκε είναι απροσδόκητα μικρό.")
            return
        except (HTTPError, URLError, TimeoutError, OSError, ValueError) as exc:
            last_error = exc
            logger.warning("Αποτυχία λήψης: %s", exc)
            if attempt < 3:
                time.sleep(attempt * 5)
    raise RuntimeError("Η λήψη του XML απέτυχε μετά από 3 προσπάθειες.") from last_error


def parse_xml(xml_path: Path) -> list[ProductRecord]:
    records: list[ProductRecord] = []
    for _event, element in ET.iterparse(xml_path, events=("end",)):
        if element.tag.split("}")[-1] != "product":
            continue
        data = {child.tag.split("}")[-1]: normalize(child.text) for child in list(element)}
        records.append(
            ProductRecord(
                code=data.get("Code", ""),
                alternative_code=data.get("Alternativecode", ""),
                barcode=data.get("Barcode", ""),
                title=data.get("Title", ""),
                stock=data.get("Stock", ""),
                wholesale_no_vat=decimal_or_none(data.get("WholesalepriceNoVat")),
                wholesale_with_vat=decimal_or_none(data.get("WholesalepricewithVat")),
                weight=decimal_or_none(data.get("Weight")),
                image=data.get("Image", ""),
            )
        )
        element.clear()
    if not records:
        raise ValueError("Δεν βρέθηκαν προϊόντα μέσα στο XML.")
    return records


def build_indexes(records: list[ProductRecord]) -> tuple[dict[str, list[ProductRecord]], dict[str, list[ProductRecord]]]:
    by_alt: dict[str, list[ProductRecord]] = {}
    by_barcode: dict[str, list[ProductRecord]] = {}
    for record in records:
        if record.alternative_code:
            by_alt.setdefault(record.alternative_code, []).append(record)
        if record.barcode:
            by_barcode.setdefault(record.barcode, []).append(record)
    return by_alt, by_barcode


def match_product(
    mpn: object,
    ean: object,
    by_alt: dict[str, list[ProductRecord]],
    by_barcode: dict[str, list[ProductRecord]],
) -> MatchResult:
    mpn_matches = by_alt.get(normalize(mpn), [])
    ean_matches = by_barcode.get(normalize(ean), [])
    if mpn_matches and ean_matches:
        intersection = [record for record in mpn_matches if record in ean_matches]
        if len(intersection) == 1:
            return MatchResult("MPN και EAN", intersection[0])
        if not intersection:
            return MatchResult("Σύγκρουση MPN/EAN", None)
        return MatchResult("Διπλή αντιστοίχιση", None)
    if len(mpn_matches) == 1:
        return MatchResult("MPN", mpn_matches[0])
    if len(ean_matches) == 1:
        return MatchResult("EAN", ean_matches[0])
    if mpn_matches or ean_matches:
        return MatchResult("Διπλή αντιστοίχιση", None)
    return MatchResult("Χωρίς αντιστοίχιση", None)


def header_map(sheet) -> dict[str, int]:
    aliases = {"ΕΑΝ": "EAN"}
    result: dict[str, int] = {}
    for cell in sheet[1]:
        label = normalize(cell.value)
        if label:
            result[aliases.get(label, label)] = cell.column
    return result


def as_float(value: object) -> float | None:
    if value is None or value == "" or isinstance(value, str) and value.startswith("="):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def calculate_profit(sheet, row: int, cost: Decimal | None, headers: dict[str, int]) -> float | None:
    if cost is None:
        return None
    net_sale = as_float(sheet.cell(row, headers["Τιμή χωρίς ΦΠΑ"]).value)
    marketplace_fee = as_float(sheet.cell(row, headers["Προμήθεια Marketplace €"]).value) or 0.0
    ad_percent = as_float(sheet.cell(row, headers["Τρέχουσα Προμήθεια Διαφήμισης %"]).value) or 0.0
    if net_sale is None:
        return None
    ad_fee = net_sale * ad_percent / 100.0
    return net_sale - float(cost) - marketplace_fee - ad_fee


def copy_cell_style(source, destination) -> None:
    from copy import copy

    if source.has_style:
        destination._style = copy(source._style)
    destination.number_format = source.number_format
    destination.alignment = copy(source.alignment)
    destination.protection = copy(source.protection)


def ensure_extra_columns(sheet) -> dict[str, int]:
    headers = header_map(sheet)
    next_column = sheet.max_column + 1
    for label in EXTRA_HEADERS:
        if label not in headers:
            headers[label] = next_column
            copy_cell_style(sheet.cell(1, 24), sheet.cell(1, next_column))
            sheet.cell(1, next_column).value = label
            next_column += 1
    widths = {
        "XML Κατάσταση Αντιστοίχισης": 25,
        "XML Κωδικός": 15,
        "XML Διαθεσιμότητα": 19,
        "XML Τιμή Χονδρικής χωρίς ΦΠΑ": 25,
        "XML Τιμή Χονδρικής με ΦΠΑ": 23,
        "XML Βάρος": 12,
        "XML Εικόνα": 42,
        "Τελευταίος Συγχρονισμός": 23,
        "Έλεγχος": 30,
    }
    for label, width in widths.items():
        sheet.column_dimensions[get_column_letter(headers[label])].width = width
        header_cell = sheet.cell(1, headers[label])
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = max(sheet.row_dimensions[1].height or 15, 42)
    return headers


def rebuild_summary(workbook, stats: dict[str, int], exceptions: list[dict[str, object]], synced_at: datetime) -> None:
    if SUMMARY_SHEET in workbook.sheetnames:
        workbook.remove(workbook[SUMMARY_SHEET])
    sheet = workbook.create_sheet(SUMMARY_SHEET, 0)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A15"
    sheet["A2"] = "Έλεγχος συγχρονισμού XML"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    sheet["A3"] = "Τελευταίος συγχρονισμός"
    sheet["B3"] = synced_at.replace(tzinfo=None)
    sheet["B3"].number_format = "dd/mm/yyyy hh:mm"
    sheet["A4"] = "Πηγή"
    sheet["B4"] = "b2bpapatheodosiou.gr XML feed (το ClientGID δεν εμφανίζεται)"

    metrics = [
        ("Προϊόντα XML", stats["xml_products"]),
        ("Προϊόντα Excel", stats["excel_products"]),
        ("Αντιστοιχισμένα", stats["matched"]),
        ("Διαθέσιμα", stats["available"]),
        ("Μη διαθέσιμα", stats["unavailable"]),
        ("Χωρίς αντιστοίχιση", stats["unmatched"]),
        ("Σύγκρουση ή διπλό", stats["ambiguous"]),
        ("Πιθανό αρνητικό κέρδος", stats["negative_profit"]),
    ]
    for index, (label, value) in enumerate(metrics, start=6):
        sheet.cell(index, 1, label)
        sheet.cell(index, 2, value)
    for row in range(6, 14):
        for column in range(1, 3):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor="F3F4F6")
            sheet.cell(row, column).font = Font(name="Arial", size=10, bold=column == 1)

    exception_headers = ["UID", "EAN", "MPN", "Όνομα", "Αντιστοίχιση", "XML Stock", "Κόστος χωρίς ΦΠΑ", "Υπολογισμένο κέρδος", "Θέμα"]
    for column, label in enumerate(exception_headers, start=1):
        cell = sheet.cell(14, column, label)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, item in enumerate(exceptions, start=15):
        values = [
            item["uid"], item["ean"], item["mpn"], item["name"], item["match_status"],
            item["stock"], item["cost"], item["profit"], item["issue"],
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        sheet.cell(row_index, 1).number_format = "@"
        sheet.cell(row_index, 2).number_format = "@"
        sheet.cell(row_index, 3).number_format = "@"
        sheet.cell(row_index, 7).number_format = '#,##0.00 "€"'
        sheet.cell(row_index, 8).number_format = '#,##0.00 "€"'
    sheet.auto_filter.ref = f"A14:I{max(14, 14 + len(exceptions))}"
    widths = [30, 18, 22, 58, 23, 19, 20, 21, 34]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[14].height = 28
    if exceptions:
        red_fill = PatternFill("solid", fgColor="FCE8E6")
        sheet.conditional_formatting.add(
            f"I15:I{14 + len(exceptions)}",
            FormulaRule(formula=["LEN(I15)>0"], fill=red_fill, font=Font(color="9C0006", bold=True)),
        )


def export_xml_products(records: list[ProductRecord], output_dir: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Προϊόντα XML"
    sheet.append(XML_EXPORT_HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for record in records:
        sheet.append([
            record.code,
            record.alternative_code,
            record.barcode,
            record.title,
            record.stock,
            float(record.wholesale_no_vat) if record.wholesale_no_vat is not None else None,
            float(record.wholesale_with_vat) if record.wholesale_with_vat is not None else None,
            float(record.weight) if record.weight is not None else None,
            record.image,
        ])

    widths = [16, 24, 20, 70, 18, 26, 25, 12, 55]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "@"
        row[1].number_format = "@"
        row[2].number_format = "@"
        row[5].number_format = '#,##0.00 "€"'
        row[6].number_format = '#,##0.00 "€"'
        row[7].number_format = "0.000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / XML_PRODUCTS_FILENAME
    workbook.save(target)
    return target


def style_audit_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [24, 14, 58, 32, 22, 18, 22, 56, 24, 16, 19, 56]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "@"
        row[1].number_format = "@"
        row[5].number_format = "@"
        row[6].number_format = "@"
        row[7].number_format = "@"
        row[8].number_format = "@"
        row[9].number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34


def build_product_audit_row(sheet, row: int, headers: dict[str, int], result: MatchResult) -> list[object]:
    product = result.product
    return [
        normalize(sheet.cell(row, headers["UID"]).value),
        normalize(sheet.cell(row, headers.get("SKU ID", 0)).value) if headers.get("SKU ID") else "",
        normalize(sheet.cell(row, headers["Όνομα"]).value),
        normalize(sheet.cell(row, headers.get("Κατηγορία προϊόντος", 0)).value) if headers.get("Κατηγορία προϊόντος") else "",
        normalize(sheet.cell(row, headers.get("Κατασκευαστής", 0)).value) if headers.get("Κατασκευαστής") else "",
        normalize(sheet.cell(row, headers["EAN"]).value),
        normalize(sheet.cell(row, headers["MPN"]).value),
        normalize(sheet.cell(row, headers.get("URL προϊόντος", 0)).value) if headers.get("URL προϊόντος") else "",
        result.status,
        product.code if product else "",
        product.stock if product else "",
        product.image if product else "",
    ]


def export_product_audit(source_path: Path, output_dir: Path, worksheet_name: str, records: list[ProductRecord]) -> tuple[Path, dict[str, int]]:
    if not source_path.exists():
        raise FileNotFoundError(f"Δεν βρέθηκε το Excel: {source_path}")
    source_workbook = load_workbook(source_path, data_only=False)
    if worksheet_name not in source_workbook.sheetnames:
        raise KeyError(f"Δεν βρέθηκε το φύλλο '{worksheet_name}'.")
    source_sheet = source_workbook[worksheet_name]
    headers = header_map(source_sheet)
    required = ["UID", "Όνομα", "EAN", "MPN"]
    missing = [label for label in required if label not in headers]
    if missing:
        raise ValueError("Λείπουν απαιτούμενες στήλες: " + ", ".join(missing))

    by_alt, by_barcode = build_indexes(records)
    unmatched_rows: list[list[object]] = []
    no_image_rows: list[list[object]] = []
    checked = 0
    matched = 0

    for row in range(2, source_sheet.max_row + 1):
        uid = normalize(source_sheet.cell(row, headers["UID"]).value)
        name = normalize(source_sheet.cell(row, headers["Όνομα"]).value)
        if not uid or not name:
            continue
        checked += 1
        result = match_product(
            source_sheet.cell(row, headers["MPN"]).value,
            source_sheet.cell(row, headers["EAN"]).value,
            by_alt,
            by_barcode,
        )
        product = result.product
        if not product:
            unmatched_rows.append(build_product_audit_row(source_sheet, row, headers, result))
            continue
        matched += 1
        if not normalize(product.image):
            no_image_rows.append(build_product_audit_row(source_sheet, row, headers, result))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Σύνοψη"
    summary_rows = [
        ("Έλεγχος", "Προϊόντα Skroutz vs XML"),
        ("Προϊόντα Excel που ελέγχθηκαν", checked),
        ("Προϊόντα XML", len(records)),
        ("Αντιστοιχισμένα", matched),
        ("Χωρίς αντιστοίχιση", len(unmatched_rows)),
        ("Χωρίς φωτογραφία XML", len(no_image_rows)),
        ("Δημιουργήθηκε", datetime.now().astimezone().replace(tzinfo=None)),
    ]
    for index, (label, value) in enumerate(summary_rows, start=1):
        summary.cell(index, 1, label)
        summary.cell(index, 2, value)
    summary["A1"].font = Font(name="Arial", size=13, bold=True)
    summary["B7"].number_format = "dd/mm/yyyy hh:mm"
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 22
    summary.sheet_view.showGridLines = False

    unmatched_sheet = workbook.create_sheet("Χωρίς αντιστοίχιση")
    unmatched_sheet.append(PRODUCT_AUDIT_HEADERS)
    for item in unmatched_rows:
        unmatched_sheet.append(item)
    style_audit_sheet(unmatched_sheet)

    no_image_sheet = workbook.create_sheet("Χωρίς φωτογραφία XML")
    no_image_sheet.append(PRODUCT_AUDIT_HEADERS)
    for item in no_image_rows:
        no_image_sheet.append(item)
    style_audit_sheet(no_image_sheet)

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / PRODUCT_AUDIT_FILENAME
    workbook.save(target)
    return target, {
        "checked": checked,
        "xml_products": len(records),
        "matched": matched,
        "unmatched": len(unmatched_rows),
        "no_image": len(no_image_rows),
    }


def update_workbook(
    source_path: Path,
    output_dir: Path,
    worksheet_name: str,
    records: list[ProductRecord],
    logger: logging.Logger,
) -> tuple[Path, dict[str, int]]:
    if not source_path.exists():
        raise FileNotFoundError(f"Δεν βρέθηκε το Excel: {source_path}")
    workbook = load_workbook(source_path, data_only=False)
    if worksheet_name not in workbook.sheetnames:
        raise KeyError(f"Δεν βρέθηκε το φύλλο '{worksheet_name}'.")
    sheet = workbook[worksheet_name]
    headers = header_map(sheet)
    required = ["UID", "Όνομα", "EAN", "MPN", "Τιμή χονδρικής", "Τιμή χωρίς ΦΠΑ", "Προμήθεια Marketplace €", "Τρέχουσα Προμήθεια Διαφήμισης %"]
    missing = [label for label in required if label not in headers]
    if missing:
        raise ValueError("Λείπουν απαιτούμενες στήλες: " + ", ".join(missing))
    headers.update(ensure_extra_columns(sheet))
    by_alt, by_barcode = build_indexes(records)
    synced_at = datetime.now().astimezone()
    stats = {
        "xml_products": len(records), "excel_products": 0, "matched": 0,
        "available": 0, "unavailable": 0, "unmatched": 0,
        "ambiguous": 0, "negative_profit": 0,
    }
    exceptions: list[dict[str, object]] = []

    for row in range(2, sheet.max_row + 1):
        uid = normalize(sheet.cell(row, headers["UID"]).value)
        name = normalize(sheet.cell(row, headers["Όνομα"]).value)
        ean = normalize(sheet.cell(row, headers["EAN"]).value)
        mpn = normalize(sheet.cell(row, headers["MPN"]).value)
        # The export ends with explanatory notes in column A. A product row
        # always has both its UID and its product name.
        if not uid or not name:
            continue
        stats["excel_products"] += 1
        result = match_product(mpn, ean, by_alt, by_barcode)
        product = result.product
        issue_parts: list[str] = []
        profit: float | None = None
        if product:
            stats["matched"] += 1
            if product.stock.casefold() == "διαθέσιμο".casefold():
                stats["available"] += 1
            else:
                stats["unavailable"] += 1
                issue_parts.append("Μη διαθέσιμο στο XML")
            if product.wholesale_no_vat is not None:
                sheet.cell(row, headers["Τιμή χονδρικής"], float(product.wholesale_no_vat))
                sheet.cell(row, headers["Τιμή χονδρικής"]).number_format = '#,##0.00 "€"'
            profit = calculate_profit(sheet, row, product.wholesale_no_vat, headers)
            if profit is not None and profit < 0:
                stats["negative_profit"] += 1
                issue_parts.append("Πιθανό αρνητικό κέρδος")
        elif result.status == "Χωρίς αντιστοίχιση":
            stats["unmatched"] += 1
            issue_parts.append("Χωρίς αντιστοίχιση")
        else:
            stats["ambiguous"] += 1
            issue_parts.append(result.status)

        values = {
            "XML Κατάσταση Αντιστοίχισης": result.status,
            "XML Κωδικός": product.code if product else "",
            "XML Διαθεσιμότητα": product.stock if product else "",
            "XML Τιμή Χονδρικής χωρίς ΦΠΑ": float(product.wholesale_no_vat) if product and product.wholesale_no_vat is not None else None,
            "XML Τιμή Χονδρικής με ΦΠΑ": float(product.wholesale_with_vat) if product and product.wholesale_with_vat is not None else None,
            "XML Βάρος": float(product.weight) if product and product.weight is not None else None,
            "XML Εικόνα": product.image if product else "",
            "Τελευταίος Συγχρονισμός": synced_at.replace(tzinfo=None),
            "Έλεγχος": "; ".join(issue_parts) if issue_parts else "OK",
        }
        for label, value in values.items():
            destination = sheet.cell(row, headers[label])
            copy_cell_style(sheet.cell(row, 24), destination)
            destination.value = value
        for label in ("XML Τιμή Χονδρικής χωρίς ΦΠΑ", "XML Τιμή Χονδρικής με ΦΠΑ"):
            sheet.cell(row, headers[label]).number_format = '#,##0.00 "€"'
        for label in ("XML Κατάσταση Αντιστοίχισης", "XML Κωδικός", "XML Διαθεσιμότητα", "XML Εικόνα", "Έλεγχος"):
            sheet.cell(row, headers[label]).number_format = "@"
        sheet.cell(row, headers["XML Βάρος"]).number_format = "0.000"
        sheet.cell(row, headers["Τελευταίος Συγχρονισμός"]).number_format = "dd/mm/yyyy hh:mm"

        if issue_parts:
            exceptions.append({
                "uid": uid, "ean": ean, "mpn": mpn, "name": name,
                "match_status": result.status,
                "stock": product.stock if product else "",
                "cost": float(product.wholesale_no_vat) if product and product.wholesale_no_vat is not None else None,
                "profit": profit,
                "issue": "; ".join(issue_parts),
            })

    last_column = max(headers[label] for label in EXTRA_HEADERS)
    sheet.auto_filter.ref = f"A1:{get_column_letter(last_column)}{sheet.max_row}"
    sheet.freeze_panes = sheet.freeze_panes or "A2"
    sheet.sheet_view.showGridLines = False
    status_column = get_column_letter(headers["Έλεγχος"])
    if sheet.max_row >= 2:
        sheet.conditional_formatting.add(
            f"{status_column}2:{status_column}{sheet.max_row}",
            FormulaRule(
                formula=[f'{status_column}2<>"OK"'],
                fill=PatternFill("solid", fgColor="FCE8E6"),
                font=Font(color="9C0006", bold=True),
            ),
        )
    rebuild_summary(workbook, stats, exceptions, synced_at)
    if hasattr(workbook, "calculation"):
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / OUTPUT_FILENAME
    with tempfile.NamedTemporaryFile(prefix="nomikos_sync_", suffix=".xlsx", dir=output_dir, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        try:
            os.replace(temporary_path, target)
        except PermissionError:
            fallback = output_dir / f"Nomikos Tools - Ενημερωμένο {synced_at:%Y-%m-%d_%H%M%S}.xlsx"
            os.replace(temporary_path, fallback)
            logger.warning("Το σταθερό αρχείο ήταν ανοικτό. Αποθηκεύτηκε ως %s", fallback)
            target = fallback
    finally:
        temporary_path.unlink(missing_ok=True)

    report_path = output_dir / "τελευταία_αναφορά.csv"
    with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exceptions[0].keys()) if exceptions else ["issue"])
        writer.writeheader()
        writer.writerows(exceptions)
    return target, stats


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ωριαίος συγχρονισμός του Excel Nomikos Tools με το XML του προμηθευτή.")
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV_FILE)
    parser.add_argument("--source", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--xml-file", type=Path, help="Χρήση τοπικού XML αντί για λήψη.")
    parser.add_argument("--export-xml-products", action="store_true", help="Δημιουργεί Excel με όλα τα προϊόντα του XML και σταματά.")
    parser.add_argument("--audit-products", action="store_true", help="Δημιουργεί αναφορά για προϊόντα χωρίς αντιστοίχιση και χωρίς φωτογραφία XML.")
    parser.add_argument("--verbose", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_arguments()
    load_env(args.env_file)
    source = args.source or Path(os.environ.get("NOMIKOS_SOURCE_XLSX", ""))
    output_dir = args.output_dir or Path(os.environ.get("NOMIKOS_OUTPUT_DIR", str(DEFAULT_OUTPUT_DIR)))
    worksheet = os.environ.get("NOMIKOS_WORKSHEET", PRODUCT_SHEET)
    timeout = int(os.environ.get("NOMIKOS_TIMEOUT_SECONDS", "180"))
    logger = configure_logging(DEFAULT_LOG_DIR, args.verbose)
    if not args.export_xml_products and not str(source):
        logger.error("Δεν έχει οριστεί NOMIKOS_SOURCE_XLSX.")
        return 2
    try:
        with RunLock(DEFAULT_STATE_DIR / "sync.lock"):
            if args.xml_file:
                xml_path = args.xml_file
            else:
                url = os.environ.get("NOMIKOS_XML_URL", "")
                if not url:
                    raise ValueError("Δεν έχει οριστεί NOMIKOS_XML_URL.")
                DEFAULT_STATE_DIR.mkdir(parents=True, exist_ok=True)
                xml_path = DEFAULT_STATE_DIR / "items_download.xml"
                download_xml(url, xml_path, timeout, logger)
            records = parse_xml(xml_path)
            logger.info("Διαβάστηκαν %s προϊόντα από το XML", len(records))
            if args.export_xml_products:
                target = export_xml_products(records, output_dir)
                logger.info("Δημιουργήθηκε πλήρης κατάλογος XML: %s", target)
                print(target)
                return 0
            if args.audit_products:
                target, stats = export_product_audit(source, output_dir, worksheet, records)
                logger.info(
                    "Δημιουργήθηκε έλεγχος προϊόντων: %s | checked=%s, unmatched=%s, no_image=%s",
                    target, stats["checked"], stats["unmatched"], stats["no_image"],
                )
                print(target)
                return 0
            target, stats = update_workbook(source, output_dir, worksheet, records, logger)
            logger.info(
                "Ολοκληρώθηκε: %s | Excel=%s, matched=%s, unavailable=%s, unmatched=%s, ambiguous=%s",
                target, stats["excel_products"], stats["matched"], stats["unavailable"],
                stats["unmatched"], stats["ambiguous"],
            )
            print(target)
            return 0
    except Exception:
        logger.exception("Ο συγχρονισμός απέτυχε")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
