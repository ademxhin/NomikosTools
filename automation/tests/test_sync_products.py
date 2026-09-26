import sys
from pathlib import Path
import unittest
from decimal import Decimal
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook

from sync_products import ProductRecord, build_indexes, export_product_audit, export_xml_products, match_product, normalize


def product(code, alternative, barcode):
    return ProductRecord(code, alternative, barcode, "Test", "Διαθέσιμο", Decimal("1"), Decimal("1.24"), None, "")


class MatchingTests(unittest.TestCase):
    def setUp(self):
        self.first = product("1", "MPN-1", "111")
        self.second = product("2", "MPN-2", "222")
        self.by_alt, self.by_barcode = build_indexes([self.first, self.second])

    def test_normalizes_integer_float(self):
        self.assertEqual(normalize(123.0), "123")

    def test_prefers_agreeing_mpn_and_ean(self):
        result = match_product("MPN-1", "111", self.by_alt, self.by_barcode)
        self.assertEqual(result.status, "MPN και EAN")
        self.assertEqual(result.product, self.first)

    def test_uses_mpn_when_ean_is_missing(self):
        result = match_product("MPN-2", "", self.by_alt, self.by_barcode)
        self.assertEqual(result.status, "MPN")
        self.assertEqual(result.product, self.second)

    def test_rejects_conflicting_identifiers(self):
        result = match_product("MPN-1", "222", self.by_alt, self.by_barcode)
        self.assertEqual(result.status, "Σύγκρουση MPN/EAN")
        self.assertIsNone(result.product)

    def test_exports_xml_products_workbook(self):
        with TemporaryDirectory() as temporary:
            target = export_xml_products([self.first, self.second], Path(temporary))
            workbook = load_workbook(target)
            sheet = workbook["Προϊόντα XML"]
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A1"].value, "XML Κωδικός")
            self.assertEqual(sheet["B2"].value, "MPN-1")
            self.assertEqual(sheet["G3"].value, 1.24)

    def test_exports_product_audit_workbook(self):
        with TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Προϊόντα που εμφανίζονται"
            sheet.append(["UID", "SKU ID", "Όνομα", "Κατηγορία προϊόντος", "Κατασκευαστής", "ΕΑΝ", "MPN", "URL προϊόντος"])
            sheet.append(["u1", "s1", "Matched no image", "Cat", "Brand", "111", "MPN-1", "https://skroutz.example/1"])
            sheet.append(["u2", "s2", "Unmatched", "Cat", "Brand", "999", "MPN-X", "https://skroutz.example/2"])
            workbook.save(source)
            target, stats = export_product_audit(source, Path(temporary), "Προϊόντα που εμφανίζονται", [self.first, self.second])
            audit = load_workbook(target)
            self.assertEqual(stats["checked"], 2)
            self.assertEqual(stats["unmatched"], 1)
            self.assertEqual(stats["no_image"], 1)
            self.assertEqual(audit["Χωρίς αντιστοίχιση"]["H2"].value, "https://skroutz.example/2")
            self.assertEqual(audit["Χωρίς φωτογραφία XML"]["H2"].value, "https://skroutz.example/1")


if __name__ == "__main__":
    unittest.main()
